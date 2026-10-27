import pandas as pd
from openpyxl.styles import Font, Border, Side
from itertools import product
from datetime import datetime
from calendar import monthrange


def extract_last_name(full_name):
    """Извлекает фамилию из полного имени"""
    if pd.isna(full_name):
        return ""

    full_name = str(full_name).strip().upper()

    if ' ' in full_name:
        return full_name.split(' ')[0]
    else:
        return full_name


def parse_resources_and_comments(group_row):
    """
    Парсит ресурсы и комментарии из строки группы.
    Возвращает список кортежей (resource, comment)
    """
    resources = []
    comments = []

    # Парсим ресурсы
    resource_val = group_row.get('Ресурс', '')
    if pd.notna(resource_val) and str(resource_val).strip():
        resources = [r.strip() for r in str(resource_val).split(';') if r.strip()]
    else:
        resources = [None]

    # Парсим комментарии
    comment_val = group_row.get('Коментарий', '')
    if pd.notna(comment_val) and str(comment_val).strip():
        # Пробуем преобразовать в числа, если возможно
        for c in str(comment_val).split(';'):
            c = c.strip()
            # Пробуем преобразовать в int или float
            try:
                if '.' in c:
                    comments.append(float(c))
                else:
                    comments.append(int(c))
            except ValueError:
                comments.append(c)
    else:
        comments = [None]

    return list(product(resources, comments))


def create_report_for_combination(df_results_filtered, df_groups, period_start_ru, period_end_ru,
                                  resource_filter=None, comment_filter=None):
    """
    Создает отчет для конкретной комбинации ресурса и комментария
    """

    # Применяем фильтры к данным, если они заданы
    df_filtered = df_results_filtered.copy()

    if resource_filter is not None:
        if 'Ресурс' in df_filtered.columns:
            before_count = len(df_filtered)
            # Приводим к строке для сравнения
            df_filtered = df_filtered[df_filtered['Ресурс'].apply(lambda v: cell_matches_filter(v, resource_filter))]
            print(f"  После фильтра по ресурсу '{resource_filter}': {before_count} -> {len(df_filtered)} записей")
        else:
            print(f"  ВНИМАНИЕ: В файле статистики нет колонки 'Ресурс'! Фильтрация невозможна.")

    if comment_filter is not None:
        if 'Коментарий' in df_filtered.columns:
            before_count = len(df_filtered)
            # Приводим оба значения к одному типу для сравнения
            df_filtered = df_filtered[
                df_filtered['Коментарий'].apply(lambda v: cell_matches_filter(v, comment_filter))
                ]
            print(f"  После фильтра по комментарию '{comment_filter}': {before_count} -> {len(df_filtered)} записей")
        else:
            print(f"  ВНИМАНИЕ: В файле статистики нет колонки 'Коментарий'! Фильтрация невозможна.")

    # Сопоставление фамилий с группами
    lastname_to_group = {}

    for _, row in df_groups.iterrows():
        group_name = row['Наименование группы']
        participants = row['Участники']

        if pd.isna(participants) or not isinstance(participants, str):
            continue

        participants_list = [p.strip() for p in participants.split(';') if p.strip()]

        for employee in participants_list:
            lastname = extract_last_name(employee)
            if lastname and lastname not in lastname_to_group:
                lastname_to_group[lastname] = group_name

    # Сопоставление сотрудников из результатов с группами
    all_employees = df_filtered['ФИО'].unique().tolist()
    employee_to_group = {}

    for employee in all_employees:
        lastname = extract_last_name(employee)
        if lastname in lastname_to_group:
            employee_to_group[employee] = lastname_to_group[lastname]
        else:
            employee_to_group[employee] = None  # Явно помечаем, что группа не найдена

    # Определение групп
    known_employees = {emp for emp, grp in employee_to_group.items() if grp is not None}
    unknown_employees = [emp for emp in all_employees if emp not in known_employees]

    # Сбор статистики - все группы из файла
    stats = {}
    all_groups_from_file = set(lastname_to_group.values())

    for group_name in all_groups_from_file:
        group_employees = [emp for emp, grp in employee_to_group.items() if grp == group_name]

        if group_employees:
            group_data = df_filtered[df_filtered['ФИО'].isin(group_employees)]

            stats[group_name] = {
                'participants': group_employees,
                'tasks_count': len(group_data),
                'cpu_hours': group_data['Время ЦП, проц.ядер/ч'].sum() if len(group_data) > 0 else 0
            }
        else:
            # Показываем группу даже если нет сотрудников
            stats[group_name] = {
                'participants': [],
                'tasks_count': 0,
                'cpu_hours': 0
            }

    # Группа UNKNOWN
    if unknown_employees:
        unknown_data = df_filtered[df_filtered['ФИО'].isin(unknown_employees)]

        stats['НЕИЗВЕСТНЫЕ СОТРУДНИКИ'] = {
            'participants': unknown_employees,
            'tasks_count': len(unknown_data),
            'cpu_hours': unknown_data['Время ЦП, проц.ядер/ч'].sum() if len(unknown_data) > 0 else 0
        }

    # Группа ALL - ВСЕГДА добавляем, даже если нет сотрудников
    stats['ВСЕ СОТРУДНИКИ'] = {
        'participants': all_employees,
        'tasks_count': len(df_filtered),
        'cpu_hours': df_filtered['Время ЦП, проц.ядер/ч'].sum() if len(df_filtered) > 0 else 0
    }

    # Формирование результата
    result_rows = []

    # Сортируем группы для предсказуемого порядка, но "ВСЕ СОТРУДНИКИ" и "НЕИЗВЕСТНЫЕ СОТРУДНИКИ" помещаем в конец
    regular_groups = [g for g in stats.keys() if g not in ['ВСЕ СОТРУДНИКИ', 'НЕИЗВЕСТНЫЕ СОТРУДНИКИ']]
    special_groups = [g for g in stats.keys() if g in ['НЕИЗВЕСТНЫЕ СОТРУДНИКИ', 'ВСЕ СОТРУДНИКИ']]

    for group_name in sorted(regular_groups) + special_groups:
        data = stats[group_name]
        cpu = round(data['cpu_hours'], 2)
        finance = round(cpu * 0.73, 2)
        participants_str = '; '.join(data['participants']) if data['participants'] else ''
        result_rows.append({
            'Группа': group_name,
            'Участники': participants_str,
            'Количество запущенных задач': data['tasks_count'],
            'Использовано ядрочасов': cpu,
            'Затраченные финансы (Руб.)': finance
        })

    return pd.DataFrame(result_rows)

def end_of_day_if_date_only(dt):
    """
    Если пользователь ввел дату без времени (например 2025-09-30),
    делаем границу включительной до конца этого дня.
    Иначе оставляем указанное время как включительную границу.
    """
    dt = pd.to_datetime(dt)
    if dt.time() == datetime.min.time():
        return dt + pd.Timedelta(days=1) - pd.Timedelta(microseconds=1)
    return dt


def normalize_for_filter(value):
    """Нормализация значений для фильтров Ресурс/Коментарий."""
    if pd.isna(value):
        return None
    s = str(value).strip()
    if not s:
        return None
    if s.endswith('.0'):
        s = s[:-2]
    # 0600, 600 и 600.0 должны считаться одним и тем же комментарием
    if s.isdigit():
        s = str(int(s))
    return s


def split_filter_tokens(value):
    """
    Возвращает нормализованные токены из ячейки.
    В исходной статистике встречаются значения вида:
      id_user=6410\nusername=daostanko\n8888
    Старый код сравнивал такую ячейку целиком и не относил её к комментарию 8888.
    Поэтому делим не только по ';', но и по переносам строк.
    """
    if pd.isna(value):
        return []
    text = str(value).replace('\r', '\n')
    raw_tokens = []
    for part in text.split(';'):
        raw_tokens.extend(part.split('\n'))
    tokens = []
    for token in raw_tokens:
        norm = normalize_for_filter(token)
        if norm is not None:
            tokens.append(norm)
    return tokens


def cell_matches_filter(cell_value, filter_value):
    """Проверяет, что фильтр есть среди токенов ячейки, а не только равен всей ячейке."""
    target = normalize_for_filter(filter_value)
    if target is None:
        return True
    return target in split_filter_tokens(cell_value)


def split_period_into_months(start_date, end_date):
    """
    Разбивает период на месяцы.
    ВАЖНО: конец месяца выставляется на 23:59:59.999999, а не на 00:00:00,
    чтобы не пропускать строки за последний день месяца.
    Возвращает список кортежей (start_month, end_month, month_name).
    """
    months = []
    current = start_date.replace(day=1)

    while current <= end_date:
        month_start = max(current, start_date)
        last_day = monthrange(current.year, current.month)[1]
        month_end_full = datetime(current.year, current.month, last_day, 23, 59, 59, 999999)
        month_end = min(month_end_full, end_date)

        month_name = f"{current.year}-{current.month:02d}"
        months.append((month_start, month_end, month_name))

        # Переход к следующему месяцу
        if current.month == 12:
            current = current.replace(year=current.year + 1, month=1)
        else:
            current = current.replace(month=current.month + 1)

    return months


def process_report(groups_file, results_file, period_start, period_end, output_file, serial_mode=False):
    """
    Формирует отчет по использованию ресурсов в разрезе групп.
    Если serial_mode=True, разбивает период на месяцы и создает таблицы для каждого месяца на одном листе.
    """

    # Загрузка данных
    print("Загрузка файла групп...")
    df_groups = pd.read_excel(groups_file, sheet_name=0, header=0)
    print(f"  Загружено групп: {len(df_groups)}")

    print("Загрузка файла статистики...")
    df_results = pd.read_excel(results_file, sheet_name=0, header=0)
    print(f"  Загружено записей: {len(df_results)}")
    print(f"  Колонки в файле статистики: {df_results.columns.tolist()}")

    # Показываем уникальные значения ресурсов и комментариев для отладки
    if 'Ресурс' in df_results.columns:
        print(f"  Уникальные ресурсы в статистике: {df_results['Ресурс'].unique()[:10]}")
    if 'Коментарий' in df_results.columns:
        print(f"  Уникальные комментарии в статистике: {df_results['Коментарий'].unique()[:10]}")

    # Проверка наличия необходимых колонок
    if 'Время ЦП, проц.ядер/ч' not in df_results.columns:
        print(f"ОШИБКА: В файле статистики нет колонки 'Время ЦП, проц.ядер/ч'")
        return

    # Преобразование данных
    print("Преобразование данных...")
    df_results['Время ЦП, проц.ядер/ч'] = df_results['Время ЦП, проц.ядер/ч'].astype(str).str.replace(',', '.').astype(
        float)
    df_results['Начало счёта'] = pd.to_datetime(df_results['Начало счёта'])#, format='mixed', dayfirst=False

    # Фильтрация по периоду
    period_start_dt = pd.to_datetime(period_start)
    period_end_dt = end_of_day_if_date_only(period_end)

    print(f"Фильтрация по периоду: {period_start} - {period_end}")
    df_results_filtered = df_results[
        (df_results['Начало счёта'] >= period_start_dt) &
        (df_results['Начало счёта'] <= period_end_dt)
        ]
    print(f"  После фильтра по периоду: {len(df_results_filtered)} записей из {len(df_results)}")
    if len(df_results_filtered) != len(df_results):
        skipped_by_period = len(df_results) - len(df_results_filtered)
        print(f"  ВНИМАНИЕ: строк вне выбранного периода: {skipped_by_period}")

    if len(df_results_filtered) == 0:
        print(f"Нет данных за период {period_start} - {period_end}")
        empty_result = pd.DataFrame(
            columns=['Группа', 'Участники', 'Количество запущенных задач', 'Использовано ядрочасов',
                     'Затраченные финансы (Руб.)'])
        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            empty_result.to_excel(writer, sheet_name='Нет_данных', index=False)
        return

    # Формирование русских форматов дат
    period_start_ru = '.'.join(period_start.split('-')[::-1])
    period_end_ru = '.'.join(period_end.split('-')[::-1])

    # Собираем все уникальные комбинации ресурсов и комментариев из файла групп
    all_combinations = set()
    for _, row in df_groups.iterrows():
        combinations = parse_resources_and_comments(row)
        all_combinations.update(combinations)

    # Убираем комбинацию (None, None)
    all_combinations.discard((None, None))

    # Если нет комбинаций, добавляем (None, None) для стандартного отчета
    if not all_combinations:
        all_combinations.add((None, None))

    print(f"\nНайдено комбинаций ресурс+комментарий: {len(all_combinations)}")
    for res, comm in all_combinations:
        print(f"  - ресурс='{res}', комментарий='{comm}'")

    # Создаем Excel файл
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        sheet_created = False

        # Если серийный режим - разбиваем на месяцы
        if serial_mode:
            months = split_period_into_months(period_start_dt, period_end_dt)
            print(f"\n=== СЕРИЙНЫЙ РЕЖИМ ===")
            print(f"Период разбит на {len(months)} месяцев:")
            for ms, me, mn in months:
                print(f"  {mn}: {ms.date()} - {me.date()}")

            # Для каждой комбинации ресурс+комментарий создаем отдельный лист
            for resource, comment in sorted(all_combinations, key=lambda x: (str(x[0] or ''), str(x[1] or ''))):
                # Формируем название листа
                sheet_name_parts = []
                if resource and str(resource) != 'None':
                    sheet_name_parts.append(str(resource))
                if comment and str(comment) != 'None':
                    sheet_name_parts.append(str(comment))

                sheet_name = "_".join(sheet_name_parts) if sheet_name_parts else "без_фильтров"
                # Очищаем название листа от недопустимых символов
                invalid_chars = ['\\', '/', '*', '?', ':', '[', ']']
                for ch in invalid_chars:
                    sheet_name = sheet_name.replace(ch, '_')
                if len(sheet_name) > 31:
                    sheet_name = sheet_name[:28] + "..."

                print(f"\nСоздаем лист '{sheet_name}' (ресурс={resource}, комментарий={comment})")

                # Подготавливаем фильтры (преобразуем None в None)
                res_filter = resource if resource and str(resource) != 'None' else None
                comm_filter = comment if comment and str(comment) != 'None' else None

                # Список для хранения данных по месяцам
                monthly_data = []
                current_row = 0

                for month_start, month_end, month_name in months:
                    # Фильтруем данные по месяцу
                    month_start_ru = '.'.join(str(month_start.date()).split('-')[::-1])
                    month_end_ru = '.'.join(str(month_end.date()).split('-')[::-1])

                    df_month = df_results_filtered[
                        (df_results_filtered['Начало счёта'] >= month_start) &
                        (df_results_filtered['Начало счёта'] <= month_end)
                        ]

                    print(f"    Месяц {month_name}: {len(df_month)} записей в периоде")

                    if len(df_month) > 0:
                        df_result = create_report_for_combination(
                            df_month, df_groups, month_start_ru, month_end_ru,
                            res_filter, comm_filter
                        )
                    else:
                        # Пустой DataFrame для месяца без данных
                        df_result = pd.DataFrame(
                            columns=['Группа', 'Участники', 'Количество запущенных задач',
                                     'Использовано ядрочасов', 'Затраченные финансы (Руб.)']
                        )

                    if df_result is not None and len(df_result) > 0:
                        monthly_data.append({
                            'month_name': month_name,
                            'month_start_ru': month_start_ru,
                            'month_end_ru': month_end_ru,
                            'df': df_result
                        })

                if monthly_data:
                    # Записываем все месяцы на один лист
                    for idx, data in enumerate(monthly_data):
                        df_result = data['df']
                        month_name = data['month_name']
                        ms_ru = data['month_start_ru']
                        me_ru = data['month_end_ru']

                        # Для первого месяца начинаем с 1 строки, для остальных добавляем отступ
                        if idx == 0:
                            start_offset = 1
                        else:
                            # Добавляем 2 пустые строки между таблицами
                            start_offset = current_row + 3

                        worksheet_name = sheet_name

                        # Заголовок месяца (над таблицей)
                        # start_offset здесь 1-based для openpyxl.
                        # pd.to_excel startrow ниже 0-based, поэтому считаем строки отдельно.
                        if worksheet_name not in writer.sheets:
                            # Лист будет создан первым вызовом to_excel ниже.
                            pass

                        meta_rows = 2 + (1 if res_filter else 0) + (1 if comm_filter else 0)
                        table_startrow_zero_based = start_offset + meta_rows

                        
                        df_result.to_excel(writer, sheet_name=sheet_name, index=False, startrow=table_startrow_zero_based)

                        worksheet = writer.sheets[sheet_name]

                        worksheet.cell(row=start_offset + 0, column=1, value=f"📅 МЕСЯЦ: {month_name}")
                        worksheet.cell(row=start_offset + 0, column=1).font = Font(bold=True, size=12)

                        worksheet.cell(row=start_offset + 1, column=1, value="Период:")
                        worksheet.cell(row=start_offset + 1, column=1).font = Font(bold=True)
                        worksheet.cell(row=start_offset + 1, column=2, value=f"{ms_ru} - {me_ru}")

                        filter_row = start_offset + 1
                        if res_filter:
                            filter_row += 1
                            worksheet.cell(row=filter_row, column=1, value="Фильтр по ресурсу:")
                            worksheet.cell(row=filter_row, column=2, value=str(res_filter))
                            worksheet.cell(row=filter_row, column=1).font = Font(bold=True)
                        if comm_filter:
                            filter_row += 1
                            worksheet.cell(row=filter_row, column=1, value="Фильтр по комментарию:")
                            worksheet.cell(row=filter_row, column=2, value=str(comm_filter))
                            worksheet.cell(row=filter_row, column=1).font = Font(bold=True)

                        # Последняя занятая строка: мета-строки + строка шапки таблицы + строки данных.
                        current_row = table_startrow_zero_based + len(df_result) + 1

                    # Настройка ширины колонок (один раз для листа)
                    if writer.sheets.get(sheet_name):
                        worksheet = writer.sheets[sheet_name]
                        worksheet.column_dimensions['A'].width = 35
                        worksheet.column_dimensions['B'].width = 50
                        worksheet.column_dimensions['C'].width = 25
                        worksheet.column_dimensions['D'].width = 25
                        worksheet.column_dimensions['E'].width = 25

                    sheet_created = True
                else:
                    print(f"  Нет данных для этой комбинации")
        else:
            # Обычный режим (без серийного разбиения)
            for resource, comment in sorted(all_combinations, key=lambda x: (str(x[0] or ''), str(x[1] or ''))):
                # Формируем название листа
                sheet_name_parts = []
                if resource and str(resource) != 'None':
                    sheet_name_parts.append(str(resource))
                if comment and str(comment) != 'None':
                    sheet_name_parts.append(str(comment))

                sheet_name = "_".join(sheet_name_parts) if sheet_name_parts else "без_фильтров"
                # Очищаем название листа от недопустимых символов
                invalid_chars = ['\\', '/', '*', '?', ':', '[', ']']
                for ch in invalid_chars:
                    sheet_name = sheet_name.replace(ch, '_')
                if len(sheet_name) > 31:
                    sheet_name = sheet_name[:28] + "..."

                print(f"\nСоздаем лист '{sheet_name}' (ресурс={resource}, комментарий={comment})")

                res_filter = resource if resource and str(resource) != 'None' else None
                comm_filter = comment if comment and str(comment) != 'None' else None

                df_result = create_report_for_combination(
                    df_results_filtered, df_groups, period_start_ru, period_end_ru,
                    res_filter, comm_filter
                )

                if df_result is not None and len(df_result) > 0:
                    # Сохраняем на лист, начиная с 4 строки
                    df_result.to_excel(writer, sheet_name=sheet_name, index=False, startrow=3)

                    # Форматирование
                    worksheet = writer.sheets[sheet_name]

                    # Заголовок с периодом
                    worksheet['B1'] = "Отчетный период:"
                    worksheet['B1'].font = Font(bold=True)
                    worksheet['B1'].border = Border(outline=Side(style='thin'))

                    worksheet['C1'] = f"{period_start_ru} --- {period_end_ru}"
                    worksheet['C1'].font = Font(bold=True)
                    worksheet['C1'].border = Border(outline=Side(style='thin'))

                    # Информация о фильтрах
                    current_row = 1
                    if res_filter:
                        current_row += 1
                        worksheet.cell(row=current_row, column=1, value="Фильтр по ресурсу:")
                        worksheet.cell(row=current_row, column=2, value=str(res_filter))
                        worksheet.cell(row=current_row, column=1).font = Font(bold=True)
                    if comm_filter:
                        current_row += 1
                        worksheet.cell(row=current_row, column=1, value="Фильтр по комментарию:")
                        worksheet.cell(row=current_row, column=2, value=str(comm_filter))
                        worksheet.cell(row=current_row, column=1).font = Font(bold=True)

                    # Настройка ширины
                    worksheet.column_dimensions['A'].width = 30
                    worksheet.column_dimensions['B'].width = 50
                    worksheet.column_dimensions['C'].width = 25
                    worksheet.column_dimensions['D'].width = 25
                    worksheet.column_dimensions['E'].width = 25

                    sheet_created = True
                else:
                    print(f"  Не удалось создать отчет для этой комбинации")

        if not sheet_created:
            print("Не удалось создать ни одного отчета")
            empty_df = pd.DataFrame({'Сообщение': ['Нет данных для заданных параметров']})
            empty_df.to_excel(writer, sheet_name='Нет_данных', index=False)

    print(f"\nОтчет сохранен: {output_file}")


# Использование
if __name__ == "__main__":
    print("=" * 60)
    print("ГЕНЕРАЦИЯ ОТЧЕТА ПО ИСПОЛЬЗОВАНИЮ РЕСУРСОВ")
    print("=" * 60)

    # Запрашиваем режим работы
    while True:
        serial_choice = input("Включить серийный режим (разбивка по месяцам)? (да/нет): ").strip().lower()
        if serial_choice in ['да', 'нет', 'yes', 'no', 'y', 'n']:
            serial_mode = serial_choice in ['да', 'yes', 'y']
            break
        print("Пожалуйста, введите 'да' или 'нет'")

    groups_file = input("Введите имя файла групп (с расширением): ").strip()
    results_file = input("Введите имя файла статистики (с расширением): ").strip()
    output_name_file = input("Введите название файла для сохранения (без формата): ").strip()

    period_start = input("Введите дату начала периода (ГГГГ-ММ-ДД): ").strip()
    period_end = input("Введите дату окончания периода (ГГГГ-ММ-ДД): ").strip()

    process_report(
        groups_file=groups_file,
        results_file=results_file,
        period_start=period_start,
        period_end=period_end,
        output_file=output_name_file + ".xlsx",
        serial_mode=serial_mode
    )